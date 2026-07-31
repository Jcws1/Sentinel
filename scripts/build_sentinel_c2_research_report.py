from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUT = Path(__file__).resolve().parents[1] / "docs" / "SENTINEL_C2_SENSOR_VEHICLE_RESEARCH_REPORT.docx"

NAVY = "17365D"
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "1E293B"
MUTED = "64748B"
LIGHT = "E8EEF5"
PALE = "F4F6F9"
GOLD = "B78103"
RED = "9B1C1C"
WHITE = "FFFFFF"
BLACK = "000000"
USABLE_DXA = 9360
TABLE_INDENT_DXA = 120


def set_run_font(run, size=11, color=INK, bold=False, italic=False, name="Calibri"):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    run.bold = bold
    run.italic = italic


def shade_paragraph(paragraph, fill, border=None):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = p_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        p_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    if border:
        p_bdr = OxmlElement("w:pBdr")
        left = OxmlElement("w:left")
        left.set(qn("w:val"), "single")
        left.set(qn("w:sz"), "18")
        left.set(qn("w:space"), "8")
        left.set(qn("w:color"), border)
        p_bdr.append(left)
        p_pr.append(p_bdr)


def add_bottom_rule(paragraph, color=BLUE, size=12):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), str(size))
    bottom.set(qn("w:space"), "6")
    bottom.set(qn("w:color"), color)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def keep_with_next(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.append(OxmlElement("w:keepNext"))


def prevent_split(row):
    tr_pr = row._tr.get_or_add_trPr()
    tr_pr.append(OxmlElement("w:cantSplit"))


def repeat_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for tag, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_fill(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths_dxa, indent_dxa=TABLE_INDENT_DXA):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths_dxa)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent_dxa))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    grid = table._tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths_dxa:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        prevent_split(row)
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths_dxa[idx])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def set_repeat_table_header(table):
    if table.rows:
        repeat_header(table.rows[0])


def add_hyperlink(paragraph, text, url, color=BLUE, size=11):
    part = paragraph.part
    rel_id = part.relate_to(
        url,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rel_id)
    run = OxmlElement("w:r")
    r_pr = OxmlElement("w:rPr")
    r_fonts = OxmlElement("w:rFonts")
    r_fonts.set(qn("w:ascii"), "Calibri")
    r_fonts.set(qn("w:hAnsi"), "Calibri")
    r_pr.append(r_fonts)
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    r_pr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    r_pr.append(u)
    sz = OxmlElement("w:sz")
    sz.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz)
    sz_cs = OxmlElement("w:szCs")
    sz_cs.set(qn("w:val"), str(int(size * 2)))
    r_pr.append(sz_cs)
    run.append(r_pr)
    text_node = OxmlElement("w:t")
    text_node.text = text
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


def add_page_field(paragraph):
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = " PAGE "
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])


def install_numbering(doc):
    numbering = doc.part.numbering_part.element
    existing_abs = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    existing_num = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    next_abs = max(existing_abs, default=0) + 1
    next_num = max(existing_num, default=0) + 1

    def make_abstract(abs_id, fmt, text):
        abstract = OxmlElement("w:abstractNum")
        abstract.set(qn("w:abstractNumId"), str(abs_id))
        multi = OxmlElement("w:multiLevelType")
        multi.set(qn("w:val"), "singleLevel")
        abstract.append(multi)
        lvl = OxmlElement("w:lvl")
        lvl.set(qn("w:ilvl"), "0")
        start = OxmlElement("w:start")
        start.set(qn("w:val"), "1")
        lvl.append(start)
        num_fmt = OxmlElement("w:numFmt")
        num_fmt.set(qn("w:val"), fmt)
        lvl.append(num_fmt)
        lvl_text = OxmlElement("w:lvlText")
        lvl_text.set(qn("w:val"), text)
        lvl.append(lvl_text)
        suff = OxmlElement("w:suff")
        suff.set(qn("w:val"), "tab")
        lvl.append(suff)
        p_pr = OxmlElement("w:pPr")
        tabs = OxmlElement("w:tabs")
        tab = OxmlElement("w:tab")
        tab.set(qn("w:val"), "num")
        tab.set(qn("w:pos"), "720")
        tabs.append(tab)
        p_pr.append(tabs)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "720")
        ind.set(qn("w:hanging"), "360")
        p_pr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:after"), "160")
        spacing.set(qn("w:line"), "280")
        spacing.set(qn("w:lineRule"), "auto")
        p_pr.append(spacing)
        lvl.append(p_pr)
        if fmt == "bullet":
            r_pr = OxmlElement("w:rPr")
            fonts = OxmlElement("w:rFonts")
            fonts.set(qn("w:ascii"), "Symbol")
            fonts.set(qn("w:hAnsi"), "Symbol")
            r_pr.append(fonts)
            lvl.append(r_pr)
        abstract.append(lvl)
        numbering.append(abstract)

    def make_num(num_id, abs_id):
        num = OxmlElement("w:num")
        num.set(qn("w:numId"), str(num_id))
        abstract_id = OxmlElement("w:abstractNumId")
        abstract_id.set(qn("w:val"), str(abs_id))
        num.append(abstract_id)
        numbering.append(num)

    make_abstract(next_abs, "bullet", "\uf0b7")
    make_num(next_num, next_abs)
    bullet_id = next_num
    make_abstract(next_abs + 1, "decimal", "%1.")
    make_num(next_num + 1, next_abs + 1)
    return bullet_id, next_num + 1


def apply_num(paragraph, num_id):
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = p_pr.find(qn("w:numPr"))
    if num_pr is None:
        num_pr = OxmlElement("w:numPr")
        p_pr.append(num_pr)
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId")
    num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num])


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = styles["Title"]
    title.font.name = "Calibri"
    title._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    title._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    title.font.size = Pt(28)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string(NAVY)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)

    subtitle = styles["Subtitle"]
    subtitle.font.name = "Calibri"
    subtitle._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    subtitle._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    subtitle.font.size = Pt(14)
    subtitle.font.color.rgb = RGBColor.from_string(MUTED)
    subtitle.paragraph_format.space_after = Pt(18)

    specs = {
        "Heading 1": (16, BLUE, 16, 8),
        "Heading 2": (13, BLUE, 12, 6),
        "Heading 3": (12, DARK_BLUE, 8, 4),
    }
    for name, (size, color, before, after) in specs.items():
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hp.paragraph_format.space_after = Pt(0)
    r = hp.add_run("SENTINEL  /  RESEARCH BRIEF")
    set_run_font(r, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fp.paragraph_format.space_before = Pt(0)
    r = fp.add_run("Sentinel MVP architecture  |  ")
    set_run_font(r, size=8.5, color=MUTED)
    add_page_field(fp)


def add_paragraph(doc, text="", bold_prefix=None, italic=False, align=None, after=6):
    p = doc.add_paragraph()
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    if bold_prefix and text.startswith(bold_prefix):
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, bold=True)
        r2 = p.add_run(text[len(bold_prefix):])
        set_run_font(r2, italic=italic)
    else:
        r = p.add_run(text)
        set_run_font(r, italic=italic)
    return p


def add_bullet(doc, bullet_id, text, bold_prefix=None):
    p = add_paragraph(doc, text, bold_prefix=bold_prefix, after=8)
    apply_num(p, bullet_id)
    return p


def add_number(doc, number_id, text):
    p = add_paragraph(doc, text, after=8)
    apply_num(p, number_id)
    return p


def add_callout(doc, label, text, accent=BLUE, fill=PALE):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.12)
    p.paragraph_format.right_indent = Inches(0.08)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(10)
    p.paragraph_format.line_spacing = 1.15
    shade_paragraph(p, fill, accent)
    r1 = p.add_run(f"{label.upper()}  ")
    set_run_font(r1, size=10.5, color=accent, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, size=10.5, color=INK)
    return p


def add_table(doc, headers, rows, widths_dxa):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, text in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_fill(cell, LIGHT)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        set_run_font(r, size=9.3, color=NAVY, bold=True)
    for row_values in rows:
        cells = table.add_row().cells
        for i, text in enumerate(row_values):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            r = p.add_run(str(text))
            set_run_font(r, size=9.1, color=INK)
    set_table_geometry(table, widths_dxa)
    set_repeat_table_header(table)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    return table


def add_source(doc, number, title, url, note):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.28)
    p.paragraph_format.first_line_indent = Inches(-0.28)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.0
    r = p.add_run(f"[{number}] ")
    set_run_font(r, size=9.0, color=NAVY, bold=True)
    add_hyperlink(p, title, url, size=9.0)
    r2 = p.add_run(f". {note}")
    set_run_font(r2, size=9.0, color=INK)


def build():
    doc = Document()
    configure_document(doc)
    bullet_id, number_id = install_numbering(doc)
    section = doc.sections[0]
    section.different_first_page_header_footer = True

    # Cover
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(86)
    kicker = doc.add_paragraph()
    kicker.alignment = WD_ALIGN_PARAGRAPH.CENTER
    kicker.paragraph_format.space_after = Pt(16)
    r = kicker.add_run("SENTINEL / ARCHITECTURE RESEARCH")
    set_run_font(r, size=10, color=GOLD, bold=True)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    title_run = title.add_run("Vehicle and Sensor Integration")
    set_run_font(title_run, size=28, color=NAVY, bold=True)
    subtitle = doc.add_paragraph(style="Subtitle")
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("MVP protocol scope, data model, emissions risk, and combat case studies")
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(8)
    rule.paragraph_format.space_after = Pt(30)
    add_bottom_rule(rule, GOLD, 10)
    meta = doc.add_paragraph()
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    meta.paragraph_format.space_after = Pt(8)
    r = meta.add_run("Prepared for the Sentinel student development team")
    set_run_font(r, size=11, color=MUTED, bold=True)
    meta2 = doc.add_paragraph()
    meta2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = meta2.add_run("28 July 2026  |  MVP research synthesis")
    set_run_font(r, size=10, color=MUTED)
    doc.add_page_break()

    # Executive brief
    doc.add_heading("Executive brief", level=1)
    add_callout(
        doc,
        "Decision",
        "Build Sentinel around a canonical asset-and-observation model at the edge gateway. Integrate the nearest practical platform interface, simulate that interface as a separate service, and treat degraded connectivity as a normal operating state.",
    )
    add_paragraph(
        doc,
        "Connecting battlefield vehicles, ambulances, drones, and fixed sensors to a C2 platform is feasible and already proven. What is not realistic is a complete, continuously current map of every asset. Combat evidence repeatedly shows partial deployment, mixed transports, stale updates, bandwidth limits, cyber disruption, and electronic-warfare pressure.",
    )
    add_paragraph(
        doc,
        "For the MVP, Sentinel should transmit compact tracks, asset state, health, link condition, and mission events. Raw vehicle buses and high-rate payload data stay local unless a named integration requires them. The product value is a defensible, time-aware picture of what is known, how fresh it is, and how confident the system should be.",
    )
    doc.add_heading("Recommended MVP cut", level=2)
    for text in [
        "Keep HTTP/JSON and WebSocket as Sentinel's internal application contract.",
        "Use MAVLink 2 as the first real UAV adapter and NDJSON as the deterministic replay/simulation fixture.",
        "Add SAPIENT/Protobuf only for a direct counter-UAS sensor integration; add ONVIF Profile T plus RTSP only for a direct EO/IR camera integration.",
        "Model connectivity explicitly: connected, degraded, intermittent, disconnected, emissions-restricted, and suspected-compromise.",
        "Keep LiDAR, CAN, DroneCAN, ROS 2/DDS, and vendor payload buses behind the vehicle or gateway boundary.",
        "Defer Link 16, ASTERIX, CoT/TAK, MQTT, Modbus, NMEA, LoRaWAN, and generic gRPC until a named customer, sensor, or partner creates a concrete requirement.",
    ]:
        add_bullet(doc, bullet_id, text)
    add_callout(
        doc,
        "Reality check",
        "Encryption can protect content, but it cannot make a radio transmission invisible. Sentinel must support silence, delay, store-and-forward, alternate bearers, stale-track visualization, and local autonomy.",
        accent=GOLD,
        fill="FFF8E8",
    )

    doc.add_page_break()

    # 1
    doc.add_heading("1. System boundary and handshake", level=1)
    add_paragraph(
        doc,
        "The C2 platform should not speak every low-level sensor or vehicle protocol directly. The edge gateway is the translation and trust boundary. It collects local data, normalizes it, applies policy and buffering, and exposes a stable Sentinel contract over whichever transport is available.",
    )
    doc.add_heading("Target runtime topology", level=2)
    topology = [
        "VEHICLE / SENSOR SIDE",
        "Local sensors and autopilot -> vendor or vehicle interfaces -> Sentinel edge gateway",
        "",
        "TACTICAL NETWORK",
        "Protected IP bearer(s), intermittent links, store-and-forward, optional silence",
        "",
        "C2 SIDE",
        "Sentinel ingest -> validation -> track/asset fusion -> UI and mission services",
        "",
        "SIMULATION SIDE",
        "Gazebo + standalone sensor simulator -> same edge gateway contract as real hardware",
    ]
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.18)
    p.paragraph_format.right_indent = Inches(0.18)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(12)
    p.paragraph_format.line_spacing = 1.05
    shade_paragraph(p, "EEF3F8", BLUE)
    for idx, line in enumerate(topology):
        r = p.add_run(line + ("\n" if idx < len(topology) - 1 else ""))
        set_run_font(
            r,
            size=9.3,
            color=NAVY if line in {"VEHICLE / SENSOR SIDE", "TACTICAL NETWORK", "C2 SIDE", "SIMULATION SIDE"} else INK,
            bold=line in {"VEHICLE / SENSOR SIDE", "TACTICAL NETWORK", "C2 SIDE", "SIMULATION SIDE"},
            name="Consolas",
        )

    doc.add_heading("Handshake sequence", level=2)
    for text in [
        "Discover and identify. The gateway recognizes a device or simulated adapter, records manufacturer/model/firmware, and assigns a stable asset or sensor identity.",
        "Authenticate and authorize. The device or upstream gateway proves identity; Sentinel applies allowlists, role and command permissions, and certificate or key policy.",
        "Negotiate capability. The adapter reports supported observations, commands, coordinate frames, units, rates, payload schemas, and health semantics.",
        "Synchronize time and position. The gateway records clock quality, time source, coordinate reference, and uncertainty before accepting observations.",
        "Stream and acknowledge. Telemetry and observations carry sequence, observed time, received time, quality, and source provenance. Commands carry expiry and acknowledgement.",
        "Degrade safely. On link loss, the gateway buffers approved events, stops expired commands, preserves local control, and reports state once communication returns.",
    ]:
        add_number(doc, number_id, text)

    doc.add_heading("Separation of responsibilities", level=2)
    add_table(
        doc,
        ["Component", "Owns", "Does not own"],
        [
            ("Vehicle/autopilot", "Stabilization, navigation, safety interlocks, local payload control", "C2-wide fusion or common operating picture"),
            ("Sensor simulator", "Synthetic detections, noise, latency, false alarms, sensor placement", "Trusted ground truth exposure to the UI"),
            ("Edge gateway", "Adapters, normalization, identity, buffering, policy, link health", "Flight-control loops or raw-sensor visualization by default"),
            ("Sentinel C2", "Asset state, fusion, tasking workflow, confidence and staleness UI", "Vendor-specific low-level buses"),
            ("Gazebo", "Physics and vehicle/sensor truth for simulation", "Production integration contract"),
        ],
        [1800, 3780, 3780],
    )

    # 2
    doc.add_heading("2. Asset and sensor taxonomy", level=1)
    add_paragraph(
        doc,
        "Sentinel should model platforms as assets with composable components, not as a flat list of sensor icons. A UAV, ground vehicle, manned aircraft, ambulance, or fixed site shares a common operational core and then adds mission-specific profiles.",
    )
    doc.add_heading("Common asset profile", level=2)
    common = [
        ("Identity", "Asset ID, callsign, type, affiliation, trust status, software/firmware"),
        ("Kinematics", "Position, altitude, velocity, heading, orientation, covariance or accuracy"),
        ("Readiness", "Power/fuel, endurance, faults, temperature, payload availability"),
        ("Communications", "Bearer, link state, signal/quality indicator, latency, last contact, route"),
        ("Mission/autonomy", "Mode, task, route/waypoints, command status, geofence or safety state"),
        ("Data quality", "Observed time, received time, age, confidence, provenance"),
    ]
    add_table(doc, ["Component", "Minimum C2 data"], common, [2100, 7260])

    doc.add_heading("Mission sensor profiles", level=2)
    mission_rows = [
        ("EO/IR", "Video/image reference, detection, classification, line of sight, gimbal pose", "Adapter only if camera is integrated; stream metadata separately from video"),
        ("Radar", "Plot/track, range, bearing, elevation, radial velocity, track quality", "Simulate normalized tracks; SAPIENT when a real C-UAS sensor is named"),
        ("RF/ESM", "Emitter detection, frequency band, bearing, confidence, identity hypothesis", "Simulate bearings/events; avoid raw spectrum in the C2 core"),
        ("Cooperative ID", "Transponder or friendly identity, position, status", "Normalize into asset identity and trust"),
        ("Threat warning", "Alert type, bearing, severity, timestamp, confidence", "Represent as event plus affected asset"),
        ("Acoustic/UGS", "Event, bearing or zone, class hypothesis, confidence", "Optional demo modality; event-level data only"),
        ("Weather/environment", "Wind, visibility, precipitation, pressure, hazards", "Useful for route/endurance and sensor-performance context"),
    ]
    add_table(
        doc,
        ["Profile", "C2-relevant output", "MVP treatment"],
        mission_rows[:3],
        [1450, 4710, 3200],
    )
    mission_cont = doc.add_heading("Mission sensor profiles (continued)", level=2)
    mission_cont.paragraph_format.page_break_before = True
    add_table(
        doc,
        ["Profile", "C2-relevant output", "MVP treatment"],
        mission_rows[3:],
        [1450, 4710, 3200],
    )
    add_callout(
        doc,
        "LiDAR",
        "LiDAR is useful for local navigation, obstacle avoidance, terrain mapping, landing, and perception. It is usually not a top-level C2 sensor feed. Sentinel should ingest derived objects, hazards, maps, or clearance state, not the point cloud by default.",
        accent=GOLD,
        fill="FFF8E8",
    )

    doc.add_heading("Ambulance and emergency-vehicle profile", level=2)
    add_paragraph(
        doc,
        "An ambulance is a legitimate C2 asset, but the operational picture and clinical record must be separated. C2 needs location, availability, route, capacity, protection status, link state, and high-level casualty-load indicators. Patient identity, clinical notes, and detailed physiological data belong in a restricted medical system and should cross the C2 boundary only through deliberate policy.",
    )
    for text in [
        "Operational channel: position, movement, mission status, destination, seats/stretcher capacity, fuel/power, faults, and last contact.",
        "Medical summary channel: aggregate patient count, triage categories, requested capability, contamination or isolation flags, and time sensitivity.",
        "Restricted clinical channel: identifiable patient data and detailed observations using healthcare schemas such as FHIR or a jurisdiction-specific EMS standard, with stricter access and audit.",
    ]:
        add_bullet(doc, bullet_id, text)

    # 3
    doc.add_heading("3. Canonical data shape", level=1)
    add_paragraph(
        doc,
        "Every adapter should translate into a small set of stable envelopes. This keeps the UI, fusion service, and simulator independent of MAVLink, ONVIF, SAPIENT, vendor APIs, or future partner platforms.",
    )
    doc.add_heading("Core envelopes", level=2)
    add_table(
        doc,
        ["Envelope", "Required fields", "Purpose"],
        [
            ("AssetState", "assetId, observedAt, receivedAt, pose, motion, mode, readiness, linkState, source", "Friendly/platform state and its quality"),
            ("Observation", "observationId, sensorId, observedAt, geometry, measurement, classification, confidence, source", "Unassociated sensor evidence"),
            ("Track", "trackId, state vector, covariance, classification, affiliation, confidence, contributingSources", "Fused operational object"),
            ("Event", "eventId, type, severity, location/asset, observedAt, confidence, source", "Threat, fault, medical, mission, or policy alert"),
            ("Command", "commandId, targetId, issuedAt, expiresAt, parameters, authority, acknowledgement", "Safe, auditable tasking"),
            ("AdapterHealth", "adapterId, device, firmware, heartbeat, faults, queueDepth, clockQuality", "Integration and gateway observability"),
        ],
        [1450, 5020, 2890],
    )
    doc.add_heading("Freshness and uncertainty are first-class", level=2)
    for text in [
        "Store observedAt and receivedAt separately. Arrival time is not measurement time.",
        "Expose ageSeconds, source, confidence, coordinate uncertainty, and link state to the UI.",
        "Render stale assets as ghosts or uncertainty regions instead of silently freezing a green icon.",
        "Increase uncertainty with time according to the asset's last known speed, manoeuvrability, and source quality.",
        "Reject or flag impossible time jumps, coordinate discontinuities, duplicate sequence numbers, and identity conflicts.",
    ]:
        add_bullet(doc, bullet_id, text)
    add_callout(
        doc,
        "Recommended states",
        "CONNECTED, DEGRADED, INTERMITTENT, DISCONNECTED, EMISSIONS_RESTRICTED, and SUSPECTED_COMPROMISE. These states should drive UI treatment, command eligibility, and buffering policy.",
    )

    # 4
    protocol_section_heading = doc.add_heading("4. Protocol scope and adapter strategy", level=1)
    protocol_section_heading.paragraph_format.page_break_before = True
    add_paragraph(
        doc,
        "The MVP should minimize integration surface area. Protocol support is justified by a concrete adjacent system, not by theoretical completeness.",
    )
    protocol_rows = [
        ("Must", "HTTP/JSON + WebSocket", "UI/control APIs and live normalized event delivery inside the current stack"),
        ("Must", "MAVLink 2", "First UAV/autopilot integration; telemetry, mission state, command acknowledgement"),
        ("Must", "NDJSON replay", "Deterministic simulation fixtures, test capture, fault injection, demos"),
        ("If named", "SAPIENT/Protobuf", "Direct autonomous counter-UAS sensor or effector integration"),
        ("If named", "ONVIF Profile T + RTSP", "Direct EO/IR camera control and media"),
        ("If named", "Vendor/platform API", "ParallaxOS, Lattice, Picogrid, Skynode, or named sensor gateway"),
        ("If named", "CoT/TAK", "Only when TAK interoperability is an explicit demo or customer requirement"),
        ("If named", "ASTERIX", "Only for a named surveillance radar or air-traffic feed"),
        ("Defer", "Link 16", "Do not implement; use an accredited gateway if a future program requires it"),
        ("Defer", "MQTT, Modbus, NMEA, LoRaWAN", "Add as isolated adapters only when selected hardware requires them"),
        ("Local only", "CAN/J1939, DroneCAN, ROS 2/DDS", "Keep behind the vehicle/gateway boundary; export normalized state"),
    ]
    protocol_matrix_heading = doc.add_heading("Protocol scope matrix", level=2)
    add_table(
        doc,
        ["Priority", "Protocol/interface", "Use in Sentinel"],
        protocol_rows,
        [1050, 2600, 5710],
    )
    doc.add_heading("Adapter contract", level=2)
    for text in [
        "Identity: device, adapter, asset, sensor, manufacturer, model, firmware, certificate/key reference.",
        "Capabilities: observation and command types, units, frames, rates, limits, media endpoints, safety constraints.",
        "Lifecycle: discover, connect, authenticate, configure, start, health, stop, recover.",
        "Normalization: source message plus canonical envelope, with traceable conversion and no hidden unit changes.",
        "Resilience: bounded queue, back-pressure, retry policy, deduplication, sequence tracking, store-and-forward.",
        "Testability: captured fixtures, simulated faults, schema validation, time-skew and packet-loss cases.",
    ]:
        add_bullet(doc, bullet_id, text)

    # 5
    doc.add_heading("5. Real-platform benchmark", level=1)
    add_paragraph(
        doc,
        "The market confirms the architectural boundary: mature systems expose entities, tasks, tracks, video, and normalized sensor observations, while aircraft and vehicle systems retain their own real-time control and payload buses.",
    )
    add_table(
        doc,
        ["Platform", "Publicly evidenced role", "Implication for Sentinel"],
        [
            ("Anduril Lattice", "Common entity model and task-oriented integration across sensors/effectors", "Closest overlap at the C2 data-model and orchestration layer"),
            ("Picogrid", "Edge hardware/software nodes connecting heterogeneous sensors and effectors", "Closest match to the gateway-plus-C2 architecture"),
            ("Auterion Skynode", "Onboard mission computer and autopilot ecosystem with MAVLink-facing integration", "Complementary UAV layer; connect through MAVLink/API, do not replace it"),
            ("ParallaxOS", "Vendor-positioned heterogeneous sensor fusion and effector/tasking layer", "Direct overlap; avoid duplicating protocol breadth without a target integration"),
            ("SAPIENT ecosystem", "Open autonomous sensor/effector messaging for counter-UAS", "Best optional standard for a real C-UAS demo"),
            ("TAK/CoT", "Shared situational awareness and tactical data exchange", "Useful interoperability edge, not the canonical internal model"),
        ],
        [1700, 4020, 3640],
    )
    doc.add_heading("Representative sensor suites", level=2)
    for text in [
        "MQ-9 Reaper public descriptions combine multi-spectral EO/IR, laser designation/ranging, and synthetic-aperture radar. Sentinel should receive selected platform state, payload state, tracks, imagery references, and mission events rather than raw internal buses [1].",
        "F-35 mission systems demonstrate tightly integrated radar, electro-optical, electronic-warfare, navigation, and communications functions. This reinforces the need to treat the aircraft as a fused platform source, not as a bag of independent C2 sensors [2].",
        "US Army VICTORY shows how a vehicle architecture can share position, time, and data services while reducing duplicated hardware. Sentinel should sit above that type of vehicle integration layer [3].",
    ]:
        add_bullet(doc, bullet_id, text)

    doc.add_heading("Where Sentinel should differentiate", level=2)
    add_paragraph(
        doc,
        "Sentinel does not need to claim broader protocol support than established vendors. A credible student MVP can differentiate through transparent uncertainty, low-cost simulation, adapter conformance tests, cost-asymmetry analytics, and a deliberate degraded-network experience. The best integration strategy is to connect to an existing gateway or mission computer when one exists, and write a narrow adapter only where it does not.",
    )

    # 6
    risk_section_heading = doc.add_heading("6. Battlefield feasibility and emissions risk", level=1)
    add_paragraph(
        doc,
        "Vehicle connectivity is operationally valuable because it reduces position-reporting burden, supports navigation, improves coordination, and gives commanders a shared picture. It also creates detectable and attackable communications dependencies. The correct conclusion is not 'never connect'; it is 'connect selectively, with disciplined emissions and graceful degradation.'",
    )
    risk_rows = [
        ("Detection/geolocation", "An adversary can detect or locate emissions even if payloads are encrypted", "Silence mode, burst transmission, configurable rates, visible emissions state"),
        ("Jamming", "Bearer becomes slow, intermittent, or unavailable", "Multiple bearers, queues, back-pressure, stale state, local autonomy"),
        ("Spoofing/identity abuse", "False or captured source injects believable data", "Mutual authentication, allowlists, provenance, anomaly flags, revocation"),
        ("Cyber compromise", "Gateway, account, or service becomes untrusted", "Least privilege, signed updates, segmentation, audit, suspected-compromise state"),
        ("Bandwidth exhaustion", "Video or excessive telemetry crowds out essential state", "Priority classes, metadata-first video, rate control, bounded payloads"),
        ("Partial deployment", "Only some assets have compatible kits or connectivity", "Explicit coverage gaps and human/manual reports"),
    ]
    risk_matrix_heading = doc.add_heading("Risk matrix", level=2)
    add_table(
        doc,
        ["Risk", "What it means", "MVP design response"],
        risk_rows,
        [1550, 4020, 3790],
    )
    add_callout(
        doc,
        "Security principle",
        "A secure radio is not a stealth radio. Sentinel's UI should never imply that encrypted connectivity is free of tactical exposure.",
        accent=RED,
        fill="FCEBEC",
    )
    doc.add_heading("Operational posture for the MVP", level=2)
    for text in [
        "Send essential state and events at the lowest useful rate; do not default to continuous high-rate streaming.",
        "Make emissions restrictions a commander/operator policy input, not an invisible network failure.",
        "Allow assets to remain operational when disconnected; queue only information that remains useful after delay.",
        "Expire commands and require acknowledgement. Never execute a stale task merely because a link recovered.",
        "Show source, time, confidence, and communication state at every decision point.",
    ]:
        add_bullet(doc, bullet_id, text)

    # 7
    case_studies_heading = doc.add_heading("7. Combat case studies", level=1)
    doc.add_heading("Iraq: Blue Force Tracking was useful, incomplete, and bandwidth-bound", level=2)
    add_paragraph(
        doc,
        "During Operation Iraqi Freedom, FBCB2/Blue Force Tracking gave equipped units a shared friendly-force picture, supported movement through dust and poor visibility, and reduced some voice position reporting. Public after-action material also records limited fielding, mixed systems, bandwidth constraints, and update intervals that could leave a partial or stale picture. One contemporary Armor article described 1,189 BFT packages across participating units and reported early update behavior around five minutes or 800 metres [12].",
    )
    add_paragraph(
        doc,
        "The 3/2 Stryker Brigade Combat Team's Iraq experience further shows that digital tracking did not replace communications diversity. Its long movement relied on BFT alongside tactical satellite, HF, and other bearers [13]. A GAO review found multiple incompatible blue-force tracking systems during OIF, forcing improvised bridges and demonstrating that interoperability is an operational problem, not a checkbox [14].",
    )
    add_callout(
        doc,
        "Sentinel lesson",
        "Design for mixed fleets and an incomplete picture. A stale or missing asset is a first-class state, and a gateway bridge is a normal deployment component.",
    )

    doc.add_heading("Afghanistan: incremental upgrades and mission-specific digital coordination", level=2)
    add_paragraph(
        doc,
        "The US Army's Joint Capabilities Release/JBC-P upgrade in Afghanistan illustrates practical fleet rollout rather than universal adoption: an Army report described upgrades for 275 vehicles and 32 tactical operations centre kits by February 2013 [15]. Years later, a TAK-based portal reportedly supported operators during the 2021 Afghanistan evacuation, showing the value of rapidly assembled, mission-specific situational-awareness tools [16].",
    )
    add_callout(
        doc,
        "Sentinel lesson",
        "Ship adapters and kits in useful increments. Support manual/operator-sourced events and temporary partner interfaces instead of assuming every platform arrives digitally integrated.",
    )

    doc.add_heading("Ukraine: software-defined C2 at scale, under electronic attack", level=2)
    add_paragraph(
        doc,
        "Ukraine's DELTA system is publicly described as a combat-proven situational-awareness and mission-management environment used across defence-force levels. Ukrainian and NATO sources describe integration of sensors, trackers, radars, satellites, UAV and fixed-camera streams, and robotic systems, with access from ordinary computing devices [17][18]. A January 2025 announcement said DELTA Mission Control was being used by more than 300 units for roughly 900 missions per day and had supported more than 200,000 missions [19].",
    )
    add_paragraph(
        doc,
        "The same war demonstrates the fragility of communications. RUSI reporting documents heavy dependence on commercial satellite connectivity, jamming and cyber attacks, and the danger of radio-frequency geolocation [20][21][22]. Tactical systems such as Kropyva and GIS Arta are also cited in public analysis as part of a distributed digital fires and coordination ecosystem [23].",
    )
    add_callout(
        doc,
        "Sentinel lesson",
        "The realistic product is a distributed, multi-source picture that survives disruption, not a permanently synchronized cloud dashboard. Commercial and tactical bearers can be powerful, but neither availability nor concealment is guaranteed.",
    )

    doc.add_heading("Cross-case synthesis", level=2)
    add_table(
        doc,
        ["Observed pattern", "Evidence across cases", "Sentinel requirement"],
        [
            ("Partial deployment", "Different units and vehicles receive different kits and versions", "Unknown/unintegrated assets and manual reports"),
            ("Mixed bearers", "Satellite, tactical radio, cellular/commercial, and local networks coexist", "Transport-agnostic gateway and route visibility"),
            ("Stale data", "Update intervals and outages make the picture lag reality", "Age, uncertainty growth, ghosting"),
            ("Interoperability gaps", "Multiple systems require gateways or improvised bridges", "Canonical model and narrow adapters"),
            ("EW/cyber pressure", "Jamming, geolocation, service disruption, and compromise are credible", "Silence, degradation, revocation, local autonomy"),
            ("High operational value", "Navigation, coordination, targeting workflow, and reduced voice burden", "Preserve essential state and tasking under constraint"),
        ],
        [2100, 3800, 3460],
    )

    # 8
    doc.add_heading("8. MVP demonstration and acceptance criteria", level=1)
    add_paragraph(
        doc,
        "The demonstration should prove that Sentinel behaves honestly when the network and sources are imperfect. The sensor simulator must be a separate process connected through the same edge gateway contract intended for real sensors and drones.",
    )
    doc.add_heading("Three-scene demo", level=2)
    scenes = [
        ("1. Connected baseline", "Gazebo UAV, simulated radar/RF/EO events, and an ambulance/ground vehicle publish through the edge gateway. The UI shows source, freshness, confidence, readiness, and link state."),
        ("2. Jamming and partition", "Latency and packet loss rise; one asset enters emissions-restricted mode; another disconnects. Sentinel grows uncertainty, expires a command, buffers useful events, and avoids showing frozen certainty."),
        ("3. Recovery and trust fault", "The link returns and the queue drains with original observation times. A conflicting identity or impossible jump triggers suspected-compromise treatment and prevents automatic tasking."),
    ]
    for title_text, body in scenes:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(8)
        r1 = p.add_run(title_text + ". ")
        set_run_font(r1, bold=True, color=NAVY)
        r2 = p.add_run(body)
        set_run_font(r2)

    doc.add_heading("Acceptance criteria", level=2)
    criteria = [
        ("Contract parity", "Real, replay, Gazebo, and synthetic adapters produce the same validated canonical envelopes."),
        ("Truth isolation", "The UI never reads simulator ground truth directly."),
        ("Time semantics", "Observed and received times remain distinct through buffering and replay."),
        ("Degradation", "Loss, delay, reordering, duplication, time skew, and recovery are visible and testable."),
        ("Command safety", "Commands have authority, expiry, acknowledgement, and disconnected behavior."),
        ("Data minimization", "Raw point clouds, vehicle buses, and unrequested clinical data do not enter the core C2 stream."),
        ("Protocol restraint", "Only the three must-have interfaces are implemented without a named conditional integration."),
        ("Operator honesty", "Every asset/track view exposes freshness, confidence, provenance, and link state."),
    ]
    add_table(doc, ["Gate", "Pass condition"], criteria, [2000, 7360])

    implementation_heading = doc.add_heading("Implementation order", level=2)
    _, implementation_number_id = install_numbering(doc)
    for text in [
        "Freeze canonical schemas and connectivity semantics.",
        "Finish the standalone sensor simulator and NDJSON capture/replay path.",
        "Implement the edge gateway lifecycle, buffering, validation, and health telemetry.",
        "Add MAVLink 2 telemetry and command acknowledgement against a software-in-the-loop vehicle.",
        "Update the UI for staleness, uncertainty, emissions restriction, and suspected compromise.",
        "Add one conditional real-sensor adapter only after selecting the actual demo hardware.",
    ]:
        add_number(doc, implementation_number_id, text)
    add_callout(
        doc,
        "Bottom line",
        "Vehicle-connected C2 is not a fairytale. The fairytale is pretending connectivity is universal, continuous, and tactically free. Sentinel should make uncertainty and network reality part of the product.",
        accent=GOLD,
        fill="FFF8E8",
    )

    # Sources
    source_section = doc.add_section(WD_SECTION.NEW_PAGE)
    source_section.page_width = Inches(8.5)
    source_section.page_height = Inches(11)
    source_section.top_margin = Inches(1)
    source_section.bottom_margin = Inches(1)
    source_section.left_margin = Inches(1)
    source_section.right_margin = Inches(1)
    source_section.header_distance = Inches(0.492)
    source_section.footer_distance = Inches(0.492)
    source_section.header.is_linked_to_previous = True
    source_section.footer.is_linked_to_previous = True
    source_heading = doc.add_heading("Appendix A. Source register", level=1)
    add_paragraph(
        doc,
        "Public sources used to benchmark platform roles, sensor suites, interoperability, and battlefield deployments. Vendor claims are treated as descriptions of stated capability, not independent performance validation.",
        italic=True,
    )
    sources = [
        (1, "US Air Force - MQ-9 Reaper fact sheet", "https://www.af.mil/About-Us/Fact-Sheets/Display/Article/104470/mq-9-reaper/", "Public payload and mission-system overview."),
        (2, "Lockheed Martin - F-35 Mission Systems design and verification", "https://sustainability.lockheedmartin.com/content/dam/lockheed-martin/eo/documents/webt/F-35_Mission_Systems_Design_Development_and_Verification.pdf", "Integrated mission-system architecture and sensor functions."),
        (3, "US Army - VICTORY open systems architecture for fighting vehicles", "https://www.army.mil/article/210117/reusable_and_refresh_able_open_systems_architecture_for_fighting_vehicles", "Vehicle integration and shared data/service approach."),
        (4, "Auterion - Skynode peripheral integration", "https://docs.auterion.com/hardware-integration/skynode/peripherals", "Onboard mission-computer and peripheral interfaces."),
        (5, "MAVLink - Common message set", "https://mavlink.io/en/messages/common", "Public UAV telemetry and command vocabulary."),
        (6, "Anduril - Lattice entities overview", "https://developer.anduril.com/guides/entities/overview", "Public entity-model integration guidance."),
        (7, "Picogrid - Orion API getting started", "https://docs.picogrid.com/reference/getting-started-with-orion", "Public API entry point for the edge/C2 platform."),
        (8, "Picogrid - Echodyne partner integration", "https://picogrid.com/newsroom/picogrid-adds-echodyne-to-its-partner-ecosystem", "Example radar ecosystem integration."),
        (9, "UK Government - SAPIENT autonomous sensor system", "https://www.gov.uk/guidance/sapient-autonomous-sensor-system", "Open autonomous sensor/effector integration standard."),
        (10, "ONVIF - Profile T specification", "https://www.onvif.org/profiles/profile-t/", "Standardized IP video streaming and control profile."),
        (11, "MITRE - Cursor on Target overview", "https://www.mitre.org/news-insights/publication/cursor-target", "Tactical event exchange background."),
        (12, "US Army Armor - Blue Force Tracking in Operation Iraqi Freedom", "https://www.benning.army.mil/armor/eARMOR/content/issues/2003/SEP_OCT/ArmorSeptemberOctober2003web.pdf", "Contemporary field observations, scale, benefits, and constraints."),
        (13, "RAND - 3/2 Stryker Brigade Combat Team in Iraq", "https://www.rand.org/content/dam/rand/pubs/monographs/2007/RAND_MG593.pdf", "Operational movement and communications mix."),
        (14, "US GAO - Battlefield automation interoperability in Operation Iraqi Freedom", "https://www.gao.gov/assets/gao-04-547.pdf", "Incompatible tracking systems and interoperability findings."),
        (15, "US Army - Blue Force Tracking upgrade in Afghanistan", "https://www.army.mil/article/96438/army_upgrades_blue_force_tracking_in_afghanistan_to_prepare_for_new_network", "Vehicle and TOC kit rollout figures."),
        (16, "US Air Force Research Laboratory - TAK support during Afghanistan evacuation", "https://www.afrl.af.mil/News/Article-Display/Article/3360125/afrl-technology-aids-operators-during-afghanistan-evacuation/", "Mission-specific digital coordination case."),
        (17, "Ukraine Ministry of Defence - DELTA deployed across defence-force levels", "https://mod.gov.ua/en/news/the-delta-combat-system-has-been-deployed-across-all-levels-of-defence-forces-of-ukraine", "Official description of deployment and combat use."),
        (18, "NATO ACT - DELTA at CWIX", "https://www.act.nato.int/article/delta-system-cwix/", "Interoperability and multi-source integration description."),
        (19, "Ukraine Ministry of Defence - DELTA Mission Control adoption", "https://mod.gov.ua/en/news/kateryna-chernohorenko-the-mission-control-module-is-now-accessible-to-all-delta-users", "Published usage and mission figures."),
        (20, "RUSI - Preliminary lessons from Russia's war in Ukraine", "https://static.rusi.org/special-report-202207-ukraine-final-web_0.pdf", "Communications, commercial satellite dependence, and operational lessons."),
        (21, "RUSI - Jamming and cyber attacks against space services in Ukraine", "https://www.rusi.org/explore-our-research/publications/commentary/jamming-and-cyber-attacks-how-space-being-targeted-ukraine", "Jamming and cyber disruption examples."),
        (22, "RUSI - Russian communications and the world of hertz", "https://www.rusi.org/explore-our-research/publications/commentary/russian-comms-ukraine-world-hertz", "Radio-frequency detection and geolocation risk."),
        (23, "RUSI - Tactical developments in the third year of the Russo-Ukrainian War", "https://static.rusi.org/tactical-developments-third-year-russo-ukrainian-war-february-2205.pdf", "Public analysis of tactical digital systems and adaptation."),
        (24, "NATO - Electronic warfare in a contested environment", "https://www.nato.int/cps/en/natohq/news_230917.htm", "Alliance-level electronic-warfare context."),
        (25, "Ukraine Ministry of Defence - DELTA cloud and cyber architecture", "https://mod.gov.ua/en/news/dou-day-2024-conference-ukrainian-developers-presented-delta-combat-system-1", "Official discussion of DELTA development and infrastructure."),
    ]
    for source in sources:
        add_source(doc, *source)

    # Metadata and save
    props = doc.core_properties
    props.title = "Sentinel Vehicle and Sensor Integration"
    props.subject = "MVP protocol scope, data model, emissions risk, and combat case studies"
    props.author = "Sentinel Project Team"
    props.keywords = "Sentinel, C2, edge gateway, sensors, UAV, MAVLink, battlefield networking"
    props.comments = "Research synthesis for MVP architecture planning."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
