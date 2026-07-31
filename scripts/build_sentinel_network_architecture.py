from __future__ import annotations

import math
import textwrap
from pathlib import Path

from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

from build_sentinel_c2_research_report import (
    BLUE,
    DARK_BLUE,
    GOLD,
    INK,
    LIGHT,
    MUTED,
    NAVY,
    PALE,
    RED,
    add_bottom_rule,
    add_bullet,
    add_callout,
    add_hyperlink,
    add_paragraph,
    add_source,
    add_table,
    configure_document,
    install_numbering,
    set_run_font,
)


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "SENTINEL_NETWORK_ARCHITECTURE.docx"
DIAGRAM = ROOT / "docs" / "sentinel-network-architecture.png"

WHITE = "#FFFFFF"
CANVAS = "#F8FAFC"
ZONE_BLUE = "#EFF6FF"
ZONE_GREEN = "#ECFDF5"
ZONE_SLATE = "#F1F5F9"
ZONE_PURPLE = "#F5F3FF"
NAVY_HEX = "#17365D"
BLUE_HEX = "#2E74B5"
GREEN_HEX = "#0F766E"
ORANGE_HEX = "#B45309"
PURPLE_HEX = "#6D28D9"
RED_HEX = "#9B1C1C"
SLATE_HEX = "#475569"
LIGHT_BORDER = "#CBD5E1"

FONT_REGULAR = Path(r"C:\Windows\Fonts\calibri.ttf")
FONT_BOLD = Path(r"C:\Windows\Fonts\calibrib.ttf")
FONT_ITALIC = Path(r"C:\Windows\Fonts\calibrii.ttf")


def font(size: int, bold: bool = False, italic: bool = False) -> ImageFont.FreeTypeFont:
    path = FONT_BOLD if bold else FONT_ITALIC if italic else FONT_REGULAR
    return ImageFont.truetype(str(path), size=size)


def wrap_pixels(
    draw: ImageDraw.ImageDraw,
    text: str,
    face: ImageFont.FreeTypeFont,
    max_width: int,
) -> list[str]:
    lines: list[str] = []
    for paragraph in text.split("\n"):
        if not paragraph:
            lines.append("")
            continue
        current = ""
        for word in paragraph.split():
            candidate = word if not current else f"{current} {word}"
            if draw.textbbox((0, 0), candidate, font=face)[2] <= max_width:
                current = candidate
            else:
                if current:
                    lines.append(current)
                current = word
        if current:
            lines.append(current)
    return lines


def center_text(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    face: ImageFont.FreeTypeFont,
    fill: str,
    spacing: int = 8,
) -> None:
    x1, y1, x2, y2 = box
    lines = wrap_pixels(draw, text, face, x2 - x1 - 36)
    heights = [draw.textbbox((0, 0), line or " ", font=face)[3] for line in lines]
    total = sum(heights) + spacing * max(0, len(lines) - 1)
    y = y1 + ((y2 - y1) - total) / 2
    for line, height in zip(lines, heights):
        width = draw.textbbox((0, 0), line, font=face)[2]
        draw.text(((x1 + x2 - width) / 2, y), line, font=face, fill=fill)
        y += height + spacing


def rounded_box(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    fill: str,
    outline: str,
    radius: int = 28,
    width: int = 4,
) -> None:
    draw.rounded_rectangle(box, radius=radius, fill=fill, outline=outline, width=width)


def card(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    title: str,
    body: str,
    *,
    fill: str = WHITE,
    outline: str = LIGHT_BORDER,
    title_color: str = NAVY_HEX,
    note: str | None = None,
    note_color: str = RED_HEX,
) -> None:
    x1, y1, x2, y2 = box
    rounded_box(draw, box, fill, outline, radius=24, width=4)
    draw.text((x1 + 28, y1 + 22), title, font=font(34, bold=True), fill=title_color)
    body_face = font(27)
    lines = wrap_pixels(draw, body, body_face, x2 - x1 - 56)
    y = y1 + 78
    for line in lines:
        draw.text((x1 + 28, y), line, font=body_face, fill="#243247")
        y += 35
    if note:
        note_face = font(24, bold=True)
        note_lines = wrap_pixels(draw, note, note_face, x2 - x1 - 56)
        y = y2 - 30 - (len(note_lines) * 31)
        for line in note_lines:
            draw.text((x1 + 28, y), line, font=note_face, fill=note_color)
            y += 31


def arrow_head(
    draw: ImageDraw.ImageDraw,
    start: tuple[float, float],
    end: tuple[float, float],
    color: str,
    size: int = 22,
) -> None:
    angle = math.atan2(end[1] - start[1], end[0] - start[0])
    left = (
        end[0] - size * math.cos(angle - math.pi / 6),
        end[1] - size * math.sin(angle - math.pi / 6),
    )
    right = (
        end[0] - size * math.cos(angle + math.pi / 6),
        end[1] - size * math.sin(angle + math.pi / 6),
    )
    draw.polygon([end, left, right], fill=color)


def line_arrow(
    draw: ImageDraw.ImageDraw,
    start: tuple[int, int],
    end: tuple[int, int],
    color: str,
    *,
    width: int = 8,
    dashed: bool = False,
    both: bool = False,
) -> None:
    if dashed:
        dx = end[0] - start[0]
        dy = end[1] - start[1]
        length = math.hypot(dx, dy)
        ux, uy = dx / length, dy / length
        cursor = 0.0
        while cursor < length:
            stop = min(cursor + 26, length)
            a = (start[0] + ux * cursor, start[1] + uy * cursor)
            b = (start[0] + ux * stop, start[1] + uy * stop)
            draw.line([a, b], fill=color, width=width)
            cursor += 42
    else:
        draw.line([start, end], fill=color, width=width)
    arrow_head(draw, start, end, color)
    if both:
        arrow_head(draw, end, start, color)


def arrow_label(
    draw: ImageDraw.ImageDraw,
    position: tuple[int, int],
    text: str,
    color: str,
    *,
    fill: str = WHITE,
) -> None:
    face = font(23, bold=True)
    bbox = draw.textbbox((0, 0), text, font=face)
    w = bbox[2] - bbox[0] + 24
    h = bbox[3] - bbox[1] + 16
    x, y = position
    draw.rounded_rectangle((x, y, x + w, y + h), radius=10, fill=fill, outline=color, width=2)
    draw.text((x + 12, y + 5), text, font=face, fill=color)


def module_pill(
    draw: ImageDraw.ImageDraw,
    box: tuple[int, int, int, int],
    text: str,
    *,
    fill: str,
    outline: str,
) -> None:
    rounded_box(draw, box, fill, outline, radius=16, width=3)
    center_text(draw, box, text, font(25, bold=True), NAVY_HEX, spacing=4)


def build_diagram() -> None:
    width, height = 3200, 1800
    image = Image.new("RGB", (width, height), CANVAS)
    draw = ImageDraw.Draw(image)

    draw.rectangle((0, 0, width, 116), fill=NAVY_HEX)
    draw.text((76, 27), "SENTINEL MVP NETWORK ARCHITECTURE", font=font(48, bold=True), fill=WHITE)
    subtitle = "current runtime + first bench/live adapter paths"
    sw = draw.textbbox((0, 0), subtitle, font=font(30))[2]
    draw.text((width - 76 - sw, 39), subtitle, font=font(30), fill="#D7E6F7")

    zones = [
        ((60, 145, 1170, 1700), ZONE_GREEN, GREEN_HEX, "FIELD, VEHICLE & SIMULATION PRODUCERS"),
        ((1220, 145, 2070, 1700), ZONE_SLATE, NAVY_HEX, "EDGE TRUST & NORMALIZATION ZONE"),
        ((2120, 145, 3140, 1700), ZONE_BLUE, BLUE_HEX, "OPERATOR & C2 ZONE"),
    ]
    for box, fill, outline, label in zones:
        rounded_box(draw, box, fill, outline, radius=34, width=4)
        x1, y1, _, _ = box
        draw.text((x1 + 30, y1 + 22), label, font=font(27, bold=True), fill=outline)

    card(
        draw,
        (105, 215, 1125, 510),
        "UAV / DRONE BENCH & LIVE",
        "sentinel-mavlink-adapter + MAVLink Router\n"
        "ground radio/modem <-> encrypted bearer <-> companion computer / flight controller / payload",
        outline=GREEN_HEX,
        note="Manual RC, flight stabilization and failsafes remain local.",
    )
    card(
        draw,
        (105, 545, 1125, 805),
        "NAMED BATTLEFIELD SENSOR ADAPTERS",
        "SAPIENT | ONVIF Profile T + RTSP | ASTERIX | vendor API\n"
        "radar, RF/ESM, EO/IR, acoustic/UGS, weather",
        outline=GREEN_HEX,
        note="Only add an interface after the demo hardware is named.",
    )
    card(
        draw,
        (105, 840, 1125, 1075),
        "MANNED VEHICLE / AMBULANCE",
        "vehicle or telematics gateway publishes selected position, readiness, capacity and link state",
        outline=GREEN_HEX,
        note="Raw CAN/J1939 and identifiable clinical data stay behind local boundaries.",
    )

    rounded_box(draw, (105, 1115, 1125, 1648), ZONE_PURPLE, PURPLE_HEX, radius=26, width=4)
    draw.text((135, 1140), "SIMULATION & REPLAY - ISOLATED", font=font(31, bold=True), fill=PURPLE_HEX)
    module_pill(
        draw,
        (145, 1215, 570, 1378),
        "Gazebo adapter +\nPX4 SITL  :8080",
        fill=WHITE,
        outline=PURPLE_HEX,
    )
    module_pill(
        draw,
        (660, 1215, 1085, 1378),
        "sentinel-sensor-sim\n:8091",
        fill=WHITE,
        outline=PURPLE_HEX,
    )
    module_pill(
        draw,
        (395, 1480, 835, 1608),
        "NDJSON deterministic replay",
        fill=WHITE,
        outline=SLATE_HEX,
    )
    line_arrow(draw, (570, 1296), (660, 1296), PURPLE_HEX, width=7, dashed=True)
    arrow_label(draw, (513, 1393), "private truth only", PURPLE_HEX, fill=ZONE_PURPLE)

    rounded_box(draw, (1270, 250, 2020, 1475), WHITE, NAVY_HEX, radius=34, width=6)
    center_text(
        draw,
        (1310, 278, 1980, 380),
        "SENTINEL EDGE GATEWAY  :8090",
        font(40, bold=True),
        NAVY_HEX,
    )
    center_text(
        draw,
        (1330, 366, 1960, 430),
        "the only operational device ingress / egress boundary",
        font(25, italic=True),
        SLATE_HEX,
    )
    modules = [
        ("Identity, registration & source mode", "#E8EEF5"),
        ("Schema, time, frame, unit & size validation", "#F4F6F9"),
        ("Canonical AssetState / Observation / Event", "#E8EEF5"),
        ("Bounded queue, freshness & store-forward", "#F4F6F9"),
        ("Command routing, expiry, idempotency & ACK", "#FFF8E8"),
        ("Health, metrics & audit metadata", "#ECFDF5"),
    ]
    y = 460
    for label, fill in modules:
        module_pill(draw, (1335, y, 1955, y + 118), label, fill=fill, outline=LIGHT_BORDER)
        y += 145
    center_text(
        draw,
        (1325, 1350, 1965, 1440),
        "Production build accepts no Gazebo truth schema.",
        font(26, bold=True),
        RED_HEX,
    )
    card(
        draw,
        (1320, 1515, 1970, 1650),
        "EDGE SUPPORT",
        "local key store | clock quality | logs/metrics",
        fill="#FFFDF7",
        outline=f"#{GOLD}",
        title_color=ORANGE_HEX,
    )

    rounded_box(draw, (2170, 240, 3090, 910), WHITE, BLUE_HEX, radius=32, width=6)
    center_text(
        draw,
        (2210, 270, 3050, 365),
        "SENTINEL C2 + API  :3001",
        font(42, bold=True),
        NAVY_HEX,
    )
    c2_modules = [
        ((2235, 400, 2605, 540), "World model\n& fusion"),
        ((2655, 400, 3025, 540), "Policy, authority\n& tasking"),
        ((2235, 580, 2605, 720), "Audit +\nscenario state"),
        ((2655, 580, 3025, 720), "Edge fusion\npoller / client"),
    ]
    for box, label in c2_modules:
        module_pill(draw, box, label, fill=ZONE_BLUE, outline=BLUE_HEX)
    center_text(
        draw,
        (2235, 760, 3025, 860),
        "REST snapshots + WebSocket events; operator intent becomes expiring, auditable commands",
        font(26),
        SLATE_HEX,
    )

    card(
        draw,
        (2270, 975, 2990, 1185),
        "OPERATOR UI / PHONE / LAPTOP",
        "paired session | common operating picture | provenance, freshness and link state",
        fill=WHITE,
        outline=BLUE_HEX,
        title_color=BLUE_HEX,
    )
    line_arrow(draw, (2630, 910), (2630, 975), BLUE_HEX, width=8, both=True)
    arrow_label(draw, (2675, 918), "HTTPS/REST + WSS", BLUE_HEX, fill=ZONE_BLUE)

    card(
        draw,
        (2185, 1240, 3075, 1568),
        "LOCAL / READ-ONLY SUPPORT SERVICES",
        "offline basemap + PMTiles\n"
        "CDSE environment metadata / DEM\n"
        "time source, monitoring and mission export",
        fill="#FFFFFF",
        outline=SLATE_HEX,
        title_color=SLATE_HEX,
        note="No cloud dependency is required for the operational loop.",
        note_color=GREEN_HEX,
    )
    line_arrow(draw, (2630, 1240), (2630, 1185), SLATE_HEX, width=6, dashed=True)

    line_arrow(draw, (1170, 395), (1220, 395), GREEN_HEX, width=10)
    arrow_label(draw, (870, 345), "state / observations / health", GREEN_HEX, fill=ZONE_GREEN)
    line_arrow(draw, (1220, 470), (1170, 470), ORANGE_HEX, width=10)
    arrow_label(draw, (815, 487), "authorized command + ACK", ORANGE_HEX, fill=ZONE_GREEN)

    line_arrow(draw, (1170, 675), (1220, 675), GREEN_HEX, width=9, dashed=True)
    line_arrow(draw, (1220, 745), (1170, 745), ORANGE_HEX, width=9, dashed=True)
    line_arrow(draw, (1170, 955), (1220, 955), GREEN_HEX, width=9, dashed=True)
    line_arrow(draw, (1170, 1288), (1220, 1288), GREEN_HEX, width=9)
    line_arrow(draw, (1170, 1542), (1220, 1542), SLATE_HEX, width=8)

    line_arrow(draw, (2070, 470), (2120, 470), GREEN_HEX, width=10)
    arrow_label(draw, (1795, 410), "validated data + gateway receipt time", GREEN_HEX, fill=ZONE_SLATE)
    line_arrow(draw, (2120, 575), (2070, 575), ORANGE_HEX, width=10)
    arrow_label(draw, (1900, 595), "commands: auth + expiry + idempotency", ORANGE_HEX, fill=ZONE_SLATE)

    draw.rectangle((0, 1715, width, height), fill=WHITE)
    legend = [
        (BLUE_HEX, False, "operator / application API"),
        (GREEN_HEX, False, "telemetry, observations, health"),
        (ORANGE_HEX, False, "authorized commands / acknowledgements"),
        (PURPLE_HEX, True, "simulation truth only"),
        (SLATE_HEX, True, "conditional or replay path"),
        (RED_HEX, True, "restricted / local responsibility"),
    ]
    x = 70
    for color, dashed, label in legend:
        line_arrow(draw, (x, 1754), (x + 58, 1754), color, width=6, dashed=dashed)
        x += 70
        draw.text((x, 1737), label, font=font(22, bold=True), fill=SLATE_HEX)
        x += draw.textbbox((0, 0), label, font=font(22, bold=True))[2] + 54

    DIAGRAM.parent.mkdir(parents=True, exist_ok=True)
    image.save(DIAGRAM, "PNG", optimize=True, dpi=(300, 300))


def clear_paragraph(paragraph) -> None:
    for child in list(paragraph._p):
        paragraph._p.remove(child)


def set_header_label(section, left: str, right: str = "MVP / DEMO BASELINE") -> None:
    header = section.header
    header.is_linked_to_previous = False
    paragraph = header.paragraphs[0]
    clear_paragraph(paragraph)
    paragraph.paragraph_format.space_after = Pt(0)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = paragraph.add_run(left)
    set_run_font(run, size=8.5, color=MUTED, bold=True)
    run = paragraph.add_run(f"                                                        {right}")
    set_run_font(run, size=8.5, color=MUTED)


def set_picture_alt(inline_shape, title: str, description: str) -> None:
    doc_pr = inline_shape._inline.docPr
    doc_pr.set("title", title)
    doc_pr.set("descr", description)


def add_small_note(doc: Document, text: str, *, before: int = 4, after: int = 6):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(before)
    paragraph.paragraph_format.space_after = Pt(after)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=MUTED, italic=True)
    return paragraph


def add_metadata_row(doc: Document, label: str, value: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    label_run = paragraph.add_run(f"{label}: ")
    set_run_font(label_run, size=10.5, color=NAVY, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=10.5, color=INK)


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(5)
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    set_run_font(run, size=9, color=MUTED, italic=True)
    paragraph._p.get_or_add_pPr().append(OxmlElement("w:keepNext"))


def configure_section(section, *, landscape: bool = False, margin: float = 1.0) -> None:
    if landscape:
        section.orientation = WD_ORIENT.LANDSCAPE
        section.page_width = Inches(11)
        section.page_height = Inches(8.5)
    else:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    section.top_margin = Inches(margin)
    section.bottom_margin = Inches(margin)
    section.left_margin = Inches(margin)
    section.right_margin = Inches(margin)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)


def build_doc() -> None:
    build_diagram()
    doc = Document()
    configure_document(doc)
    bullet_id, number_id = install_numbering(doc)
    section = doc.sections[0]
    section.different_first_page_header_footer = True
    set_header_label(section, "SENTINEL  /  NETWORK ARCHITECTURE")

    kicker = doc.add_paragraph()
    kicker.paragraph_format.space_before = Pt(30)
    kicker.paragraph_format.space_after = Pt(5)
    run = kicker.add_run("SENTINEL / TECHNICAL DECISION")
    set_run_font(run, size=10, color=GOLD, bold=True)

    title = doc.add_paragraph()
    title.paragraph_format.space_after = Pt(6)
    run = title.add_run("Network Architecture")
    set_run_font(run, size=28, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(14)
    run = subtitle.add_run(
        "MVP and demo topology for sensors, UAVs, manned vehicles, simulation, and degraded battlefield networks"
    )
    set_run_font(run, size=13.5, color=MUTED, italic=True)

    rule = doc.add_paragraph()
    rule.paragraph_format.space_after = Pt(14)
    add_bottom_rule(rule, GOLD, 10)
    add_metadata_row(doc, "Status", "Recommended baseline grounded in the current Sentinel implementation")
    add_metadata_row(doc, "Current services", "C2/API :3001 | edge gateway :8090 | sensor simulator :8091 | Gazebo adapter :8080")
    add_metadata_row(doc, "Design preset", "Standard business brief / memo masthead")
    add_metadata_row(doc, "Evidence", "Public platform interfaces plus Iraq, Afghanistan, and Ukraine case studies")

    add_callout(
        doc,
        "Architecture decision",
        "Sentinel C2 talks to one canonical edge boundary, never directly to a sensor, autopilot, Gazebo, or raw vehicle bus. Every live, simulated, or replay producer is translated by a separate adapter into the same time-aware asset, observation, event, health, and command contracts.",
    )

    doc.add_heading("What this architecture optimizes for", level=1)
    for text in [
        "Student implementability: use the HTTP/JSON and WebSocket stack already running in Sentinel, then add only one practical UAV adapter and one named sensor adapter.",
        "Replaceability: simulation, replay, and bench hardware must be swappable by configuration without changing the UI or C2 domain model.",
        "Battlefield honesty: stale, intermittent, emissions-restricted, disconnected, and suspected-compromise are normal states, not exceptional error screens.",
        "Cost asymmetry: move compact evidence, tracks, state, and commands; do not spend bandwidth and compute transporting raw payload buses or point clouds that do not improve a decision.",
    ]:
        add_bullet(doc, bullet_id, text)

    add_callout(
        doc,
        "Security warning",
        "The current launch script binds the C2, edge gateway, and sensor simulator to all interfaces. That is acceptable only on an isolated demo LAN. On a shared network, expose the paired C2 endpoint and bind ports 8080, 8090, and 8091 to loopback or a dedicated internal interface; add TLS/mTLS or a protected tunnel before using real hardware.",
        accent=RED,
        fill="FDECEC",
    )

    landscape = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(landscape, landscape=True, margin=0.58)
    landscape.different_first_page_header_footer = False
    landscape.header.is_linked_to_previous = True
    landscape.footer.is_linked_to_previous = True
    heading = doc.add_heading("1. Recommended logical network", level=1)
    heading.paragraph_format.space_before = Pt(0)
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_after = Pt(0)
    picture = paragraph.add_run().add_picture(str(DIAGRAM), width=Inches(9.68))
    set_picture_alt(
        picture,
        "Sentinel MVP network architecture",
        "Three-zone architecture. Field devices, real adapters, Gazebo, the standalone sensor simulator, and NDJSON replay feed a single Sentinel edge gateway. The edge gateway validates and normalizes data before passing it to Sentinel C2 and the operator UI. Gazebo truth is restricted to the sensor simulator. Current local ports are C2 3001, edge 8090, sensor simulator 8091, and Gazebo adapter 8080.",
    )
    add_caption(
        doc,
        "Figure 1. Recommended Sentinel MVP topology. Solid paths are required/current; dashed paths are conditional or isolated. Simulation truth never crosses the edge boundary.",
    )

    portrait = doc.add_section(WD_SECTION.NEW_PAGE)
    configure_section(portrait, landscape=False, margin=1.0)
    portrait.different_first_page_header_footer = False
    portrait.header.is_linked_to_previous = True
    portrait.footer.is_linked_to_previous = True

    doc.add_heading("2. How to read the diagram", level=1)
    add_paragraph(
        doc,
        "The diagram separates responsibility before it separates products. The operator and C2 zone owns the common operating picture, fusion, policy, tasking, scenario state, and audit. The edge zone owns identity, validation, normalization, buffering, source health, and command routing. Field systems retain their own safety-critical control loops and proprietary buses. The simulation lane mirrors the field interface but keeps Gazebo truth private.",
    )

    doc.add_heading("Current runtime contracts", level=2)
    add_table(
        doc,
        ["Connection", "Current implementation", "Recommended network rule"],
        [
            (
                "Operator -> Sentinel C2",
                "HTTP/REST under /api/v1 plus WebSocket /api/v1/ws; production-style launch on :3001; Vite dev proxy on :5173.",
                "Only operator-facing endpoint. Require pairing now; add HTTPS and session controls off an isolated LAN.",
            ),
            (
                "C2 -> sentinel-edge",
                "HTTP snapshots/commands and WebSocket event forwarding on :8090.",
                "Allow only the C2 host and approved adapters. Do not expose to the operator LAN.",
            ),
            (
                "sensor-sim -> edge",
                "Source registration, health, and observation batches over HTTP/JSON with x-edge-token.",
                "Loopback or simulation VLAN. A static demo token is not sufficient for a shared/live network.",
            ),
            (
                "C2 -> sensor-sim",
                "Sensor CRUD is routed through edge to :8091; the UI has no direct simulator connection.",
                "Keep :8091 internal. Edge remains the only lifecycle boundary.",
            ),
            (
                "Gazebo/PX4 -> edge",
                "Vehicle handshake, state, commands, faults, and /v1/events are proxied through edge from :8080.",
                "Private simulation interface. Replace with a MAVLink adapter for bench/live mode.",
            ),
            (
                "Gazebo -> sensor-sim",
                "Read-only state/event feed plus optional visual-sensor CRUD.",
                "Simulation-truth network only. Never route or compile the truth schema into production edge.",
            ),
            (
                "NDJSON replay -> edge",
                "Deterministic observation capture/replay using the canonical observation envelope.",
                "First-class demo fallback and conformance fixture, not a special UI mode.",
            ),
        ],
        [1800, 4300, 3260],
    )
    add_small_note(
        doc,
        "Internal evidence: Sentinel source tree and launch configuration reviewed 28 July 2026. The target architecture deliberately follows the service boundaries already implemented.",
    )

    doc.add_heading("Four architectural planes", level=2)
    planes = [
        ("Operator plane.", "Browser/phone UI, pairing, visualization, source age/confidence, and human confirmation. It never owns a device connection."),
        ("C2 plane.", "World model, association/fusion, policy, mission planning, audit, and safe command intent."),
        ("Edge/adapter plane.", "Device discovery, authentication, capability mapping, normalization, rate control, buffering, acknowledgement, and source health."),
        ("Device/simulation plane.", "Flight controllers, payloads, sensor receivers, vehicle gateways, Gazebo/PX4, synthetic models, and replay producers."),
    ]
    for label, detail in planes:
        add_bullet(doc, bullet_id, f"{label} {detail}", bold_prefix=label)

    doc.add_heading("3. Data and command paths", level=1)
    doc.add_heading("Telemetry and observations move north", level=2)
    for text in [
        "A producer registers a stable source identity, mode (LIVE, SIMULATED, or REPLAY), instance, capabilities, software/firmware, and configuration revision.",
        "The adapter timestamps and transforms native data. The edge adds receipt time, validates schema, units, coordinate frame, sequence, age, size, and source authorization.",
        "The edge emits compact AssetState, Observation, Event, and health envelopes. Imagery travels by reference or a separate media path, not inside the event stream.",
        "C2 associates and fuses observations, retains provenance and uncertainty, and publishes a time-aware picture to the UI.",
    ]:
        add_bullet(doc, bullet_id, text)

    doc.add_heading("Commands move south, with authority and expiry", level=2)
    for text in [
        "The operator requests mission intent; C2 checks policy, role, target state, freshness, and any human-confirmation requirement.",
        "C2 creates an idempotent command with target identity, authority, issue time, expiry, and parameters.",
        "The edge routes only allow-listed commands to the owning adapter and records an acknowledgement.",
        "The adapter translates mission intent into MAVLink, a vendor API, SAPIENT tasking, or another named local interface. Accepted does not mean physically completed; completion state remains explicit.",
        "If the link fails or the command expires, the asset continues its local safety behavior. Sentinel must not execute a stale task merely because communications return.",
    ]:
        add_bullet(doc, bullet_id, text)
    add_callout(
        doc,
        "Handshake rule",
        "A heartbeat makes an asset visible, not commandable. Commandability follows identity, capability discovery, supported firmware, clock quality, position/estimator health, ownership, and safety release. MAVLink supports heartbeat, metadata, time synchronization, command acknowledgement, and mission upload, but MAVLink 2 signing authenticates messages without encrypting their content [1]-[6].",
        accent=GOLD,
        fill="FFF8E8",
    )

    doc.add_heading("Failure behavior is part of the network contract", level=2)
    add_table(
        doc,
        ["Failure", "Required system behavior", "Operator-visible state"],
        [
            (
                "Sensor simulator stops",
                "Edge and C2 remain available; source TTL expires; no perfect frozen detections.",
                "Source stale/degraded; affected tracks coast and uncertainty grows.",
            ),
            (
                "Gazebo truth stops",
                "sensor-sim stops producing normal observations and reports degraded input health.",
                "Truth-input warning; no fabricated continuity.",
            ),
            (
                "Bearer jams or drops",
                "Local autonomy continues; queue only time-useful events; commands expire.",
                "Intermittent/disconnected or emissions-restricted; last contact and age visible.",
            ),
            (
                "Edge restarts",
                "Sources re-register and snapshots reconcile without duplicate sensors, vehicles, or commands.",
                "Temporary degraded state followed by explicit recovery.",
            ),
            (
                "Identity/time conflict",
                "Quarantine or reject the source; stop automatic tasking.",
                "Suspected-compromise badge and an auditable reason.",
            ),
        ],
        [1900, 4300, 3160],
    )

    doc.add_heading("4. Trust boundaries and defensive controls", level=1)
    add_table(
        doc,
        ["Boundary", "MVP control", "Bench/live hardening"],
        [
            (
                "Operator -> C2",
                "Pairing code, operator headers/session, command confirmation and audit.",
                "HTTPS, stronger identity, role separation, lockout and session expiry.",
            ),
            (
                "C2 -> edge",
                "Separate process, explicit API, bounded request size and command IDs.",
                "mTLS or WireGuard/IPsec, firewall allow-list, certificate rotation and revocation.",
            ),
            (
                "Producer -> edge",
                "Source registration, producer token, schema and sequence validation, source health.",
                "Per-device credentials, signed builds, isolated adapter accounts and least privilege.",
            ),
            (
                "Edge -> field bearer",
                "Transport-agnostic application contract and visible link state.",
                "Encrypted radio/IP bearer, alternate path policy, EMCON/silence and store-forward.",
            ),
            (
                "Gazebo truth -> sensor-sim",
                "Different schema and endpoint; sensor-sim strips truth identity from observations.",
                "Separate process/network; production edge build excludes the truth contract.",
            ),
            (
                "Ambulance -> C2",
                "Operational location/readiness/capacity only.",
                "Restricted clinical system for patient identity and detailed observations; explicit policy gateway if a summary is required.",
            ),
        ],
        [1900, 3500, 3960],
    )
    add_callout(
        doc,
        "Emissions principle",
        "Encryption protects content, not the fact, time, power, direction, or approximate origin of a transmission. Sentinel therefore needs silence, burst/rate control, alternate bearers, stale-data treatment, and disconnected operation; a permanently synchronized map is not a credible battlefield assumption.",
        accent=RED,
        fill="FDECEC",
    )

    doc.add_heading("5. Why these interfaces are in the MVP", level=1)
    add_table(
        doc,
        ["Interface or component", "Why it is included", "Scope boundary"],
        [
            (
                "HTTP/JSON + WebSocket",
                "Already implemented across UI, C2, edge, sensor-sim, and Gazebo adapter; easy to inspect and test.",
                "Internal application contract. Add transport security when crossing hosts.",
            ),
            (
                "Canonical edge envelopes",
                "Prevent every C2 service from understanding every vendor or simulation protocol.",
                "AssetState, Observation, Event, Command, acknowledgement and health; retain source/time/confidence.",
            ),
            (
                "MAVLink 2 adapter",
                "Best first UAV interface for PX4/ArduPilot and a student-accessible bench handshake [1]-[6].",
                "Separate adapter process. Keep flight loops, payload buses and safety on the vehicle.",
            ),
            (
                "MAVLink Router",
                "Allows Sentinel, a pilot/GCS, and a recorder to share one stream without making the UI the link owner.",
                "Routing is not identity, encryption, or command authority.",
            ),
            (
                "NDJSON replay",
                "Deterministic demos, fault reproduction, conformance fixtures, and a fallback when hardware fails.",
                "Same envelope and provenance as live data.",
            ),
            (
                "One named sensor adapter",
                "Proves the simulator can be replaced by hardware without changing C2.",
                "Choose SAPIENT for a C-UAS feed or ONVIF/RTSP for an IP camera only after hardware selection [11][12].",
            ),
            (
                "Metadata/media split",
                "Tracks, detections and image references survive constrained links; raw video does not crowd out state.",
                "Use a separate RTSP/WebRTC/object path for pixels.",
            ),
        ],
        [1900, 4140, 3320],
    )

    doc.add_heading("6. Why other protocols and data were not added", level=1)
    add_table(
        doc,
        ["Deferred or excluded", "Reason", "Revisit trigger"],
        [
            (
                "Link 16",
                "Accredited tactical datalink integration is not a student implementation task; protocol breadth adds no MVP proof.",
                "A programme supplies an approved gateway and concrete interface.",
            ),
            (
                "Raw CAN/J1939, DroneCAN, ROS 2/DDS",
                "They belong inside the vehicle or mission computer and would couple C2 to local control and payload buses.",
                "A named vehicle gateway requires a narrow adapter above the bus.",
            ),
            (
                "Raw LiDAR point clouds",
                "High rate and primarily useful for local navigation, mapping, landing, and collision avoidance.",
                "C2 has a specific derived-object, hazard-map, or clearance-state need.",
            ),
            (
                "Full ASTERIX, CoT/TAK, MQTT, Modbus, NMEA, LoRaWAN",
                "Each is valuable in a specific ecosystem, but speculative support creates adapters with no hardware acceptance test.",
                "A named radar, partner C2, environmental sensor, or bearer requires it.",
            ),
            (
                "Generic gRPC/Protobuf rewrite",
                "The current vertical slice works over HTTP/JSON/WS. Re-platforming does not prove hardware interchangeability.",
                "Measured throughput, cross-language needs, or a selected SAPIENT integration justify it.",
            ),
            (
                "Continuous video on WebSocket",
                "Video can dominate constrained links and block more decision-relevant state.",
                "A separate, rate-controlled media service is demonstrated.",
            ),
            (
                "Cloud-only services and Kubernetes",
                "They add operational dependence and complexity before the edge contract is proven.",
                "A deployment scale or availability requirement is measured, not assumed.",
            ),
            (
                "Direct UI -> device/Gazebo",
                "It bypasses identity, audit, validation, buffering, source health and safe command ownership.",
                "Never; use C2 and edge interfaces.",
            ),
        ],
        [2000, 4260, 3100],
    )

    doc.add_heading("7. Benchmark against real platforms", level=1)
    add_paragraph(
        doc,
        "The goal is not to reproduce an established platform. It is to copy the architectural patterns that recur publicly, then stop before Sentinel becomes an untestable collection of protocol stubs. Vendor descriptions below are evidence of stated integration models, not independent performance validation.",
    )
    add_table(
        doc,
        ["Platform / standard", "Publicly visible pattern", "Design consequence for Sentinel"],
        [
            (
                "Anduril Lattice",
                "Composable entities, assets/tracks, provenance, expiry, health, real-time streams, task catalogs, REST and gRPC [8][9].",
                "Use a canonical C2 entity/observation model and robust partial/stale state; do not let native device protocols leak upward.",
            ),
            (
                "Picogrid Legion + ECNs",
                "Legion provides a common interface; Lander/Helios/Portal provide power, communications and edge compute for sensors, drones and vehicles [10].",
                "The closest public analogue to Sentinel's gateway-plus-C2 boundary; field compute is part of integration, not an afterthought.",
            ),
            (
                "Auterion Skynode",
                "Flight controller plus Linux mission computer, peripheral interfaces and MAVLink-facing integration [7].",
                "Treat the onboard mission computer as the UAV integration layer. Sentinel should connect to it, not replace flight control.",
            ),
            (
                "Surtr ParallaxOS",
                "Vendor-described hardware-agnostic C-UAS layer fusing radar, RF, camera and acoustic sources; public implementation detail remains limited [22].",
                "Direct product overlap at sensor fusion/C2. Differentiate with transparent uncertainty, deterministic simulation, conformance tests and low-cost deployment.",
            ),
            (
                "SAPIENT",
                "Open autonomous sensor/effector messaging with registration, status, detections, tasks and acknowledgements [11].",
                "Best conditional interface if the chosen demo is a real counter-UAS sensor, but not mandatory without that hardware.",
            ),
            (
                "TAK / Cursor on Target",
                "Shared situational-awareness and tactical event exchange [13].",
                "Useful interoperability edge; keep CoT as an adapter, not Sentinel's internal model.",
            ),
            (
                "US Army VICTORY",
                "Vehicle open architecture shares data and services while reducing duplicated hardware [14].",
                "Sentinel should sit above a vehicle gateway and ingest selected state rather than attach to every CAN/J1939 signal.",
            ),
        ],
        [1700, 3900, 3760],
    )

    doc.add_heading("8. Combat evidence that changes the architecture", level=1)
    doc.add_heading("Iraq: useful digital tracking, but partial and bandwidth-bound", level=2)
    add_paragraph(
        doc,
        "Blue Force Tracking gave equipped units a shared friendly-force picture and supported movement in poor visibility, but public reporting also records limited fielding, update intervals, bandwidth constraints, and incompatible systems. A contemporary Armor article described 1,189 BFT packages and early update behavior around five minutes or 800 metres; a GAO review found multiple incompatible tracking systems that required improvised bridges [15][16].",
    )
    add_callout(
        doc,
        "Sentinel consequence",
        "Unknown, manual, stale, and unintegrated assets are first-class states. A gateway bridge and mixed software versions are normal deployment conditions.",
    )

    doc.add_heading("Afghanistan: incremental kits and mission-specific coordination", level=2)
    add_paragraph(
        doc,
        "The US Army publicly described a JBC-P/Blue Force Tracking upgrade covering 275 vehicles and 32 tactical-operations-centre kits by February 2013. AFRL later described TAK-based support during the 2021 evacuation [17][23]. These are evidence for useful incremental rollout and rapidly assembled coordination, not universal native integration.",
    )
    add_callout(
        doc,
        "Sentinel consequence",
        "Ship adapters and edge kits in useful increments. Preserve manual/operator-sourced events and temporary partner interfaces.",
    )

    doc.add_heading("Ukraine: software-defined C2 under electronic attack", level=2)
    add_paragraph(
        doc,
        "Ukraine and NATO publicly describe DELTA as a multi-source situational-awareness and mission-management environment integrating sensors, trackers, radars, satellites, UAVs, cameras, and robotic systems. A January 2025 announcement reported more than 300 units, roughly 900 missions per day, and more than 200,000 supported missions [18][19]. The same war demonstrates dependence on mixed commercial and tactical bearers, jamming, cyber disruption, and radio-frequency geolocation risk [20][21].",
    )
    add_callout(
        doc,
        "Sentinel consequence",
        "Build a distributed, multi-source picture that survives disruption. Connectivity is valuable, but availability and concealment are never guaranteed.",
    )

    doc.add_heading("Cross-case architecture requirements", level=2)
    for text in [
        "Partial deployment: represent assets with unknown or manual-source status rather than silently deleting them.",
        "Mixed bearers: make transport and route health visible while keeping the application contract bearer-agnostic.",
        "Stale data: store observation time and receipt time separately; grow uncertainty instead of freezing a green icon.",
        "Interoperability gaps: use narrow adapters around a canonical model and expect temporary bridges.",
        "EW/cyber pressure: support silence, revocation, suspected-compromise, local autonomy, and explicit recovery.",
    ]:
        add_bullet(doc, bullet_id, text)

    doc.add_heading("9. Minimum hardware topology for the first bench demo", level=1)
    add_table(
        doc,
        ["Location", "Minimum hardware/software", "Responsibility"],
        [
            (
                "On the UAV",
                "Pixhawk-class flight controller; PX4 or ArduPilot; GPS/IMU; air telemetry radio or IP modem; antenna; optional Linux companion computer; independent manual-control/failsafe path.",
                "Stable flight, estimator, safety and local payload control.",
            ),
            (
                "Ground link",
                "Matching ground radio/modem; USB/serial/Ethernet; MAVLink Router; sentinel-mavlink-adapter.",
                "Physical bearer, stream fan-out, identity/capability mapping and command acknowledgement.",
            ),
            (
                "Edge host",
                "Student laptop or small industrial PC; sentinel-edge; key store; bounded recorder; health/metrics.",
                "Only canonical device ingress/egress; can be co-located for the demo.",
            ),
            (
                "C2/operator host",
                "Sentinel C2/API and UI; offline map data; scenario/audit store.",
                "Fusion, policy, tasking, visualization and human confirmation.",
            ),
            (
                "First real sensor",
                "Named radar or camera plus its required receiver/SDK; one adapter selected from SAPIENT, ONVIF/RTSP, ASTERIX, or vendor API.",
                "Prove live/sim interchangeability with a real acceptance fixture.",
            ),
            (
                "Ambulance/vehicle option",
                "Telematics or vehicle gateway exposing position/readiness/capacity; separate restricted medical endpoint if needed.",
                "Operational overview without pulling raw vehicle or patient records into C2.",
            ),
        ],
        [1600, 4550, 3210],
    )
    add_callout(
        doc,
        "Radio reality",
        "Pairing the telemetry radios or establishing an IP bearer happens before MAVLink. MAVLink is the application protocol carried by serial, UDP/TCP, Ethernet, Wi-Fi, LTE, mesh, or another radio path; it does not replace the modem, antenna, ground receiver, or encrypted transport.",
        accent=GOLD,
        fill="FFF8E8",
    )

    doc.add_heading("Recommended demo firewall posture", level=2)
    for text in [
        "Operator LAN -> allow TCP :3001 only; require pairing and prefer HTTPS.",
        "C2 host -> allow edge :8090; deny direct access to sensor-sim :8091 and Gazebo adapter :8080 from operator devices.",
        "Adapter/simulation network -> allow only the expected producer and management paths; block internet ingress.",
        "If Gazebo runs in WSL or another host, use a dedicated host-only/isolated link and never route the truth endpoint through the operator Wi-Fi.",
        "Record source, last contact, link state, queue depth and dropped messages so degraded behavior is visible during the demo.",
    ]:
        add_bullet(doc, bullet_id, text)

    doc.add_heading("10. Implementation sequence", level=1)
    sequence = [
        "Freeze the canonical schemas and connectivity states already used by the edge vertical slice.",
        "Harden process bindings: expose :3001 to the operator LAN; keep :8080/:8090/:8091 internal; remove the default producer token for shared networks.",
        "Add sentinel-mavlink-adapter as a separate process with MAVLink Router, heartbeat/capability discovery, time synchronization, signed-message support, command expiry and acknowledgement.",
        "Prove software-in-the-loop first, then repeat the same conformance tests with one physical PX4/ArduPilot vehicle.",
        "Choose one real external sensor. Implement only its SAPIENT, ONVIF/RTSP, ASTERIX, or vendor adapter and preserve raw captures for regression fixtures.",
        "Exercise loss, delay, time skew, duplicate IDs, restart, emissions restriction and suspected-compromise in the standalone sensor simulator and NDJSON replay.",
        "Only after measurement shows a bottleneck should the team consider gRPC/Protobuf transport, a media service, or wider protocol support.",
    ]
    for text in sequence:
        paragraph = add_paragraph(doc, text, after=8)
        from build_sentinel_c2_research_report import apply_num

        apply_num(paragraph, number_id)

    doc.add_heading("Architecture acceptance criteria", level=2)
    add_table(
        doc,
        ["Gate", "Pass condition"],
        [
            (
                "Single boundary",
                "The UI has no direct socket to MAVLink, a sensor, Gazebo, sensor-sim, or a raw vehicle bus.",
            ),
            (
                "Contract parity",
                "Simulation, replay, SITL, and one bench/live adapter produce the same validated canonical envelopes.",
            ),
            (
                "Truth isolation",
                "No normal edge packet contains a Gazebo entity ID or perfect target pose.",
            ),
            (
                "Time semantics",
                "Observed, sent, received and gateway-receipt times remain distinct through buffering and replay.",
            ),
            (
                "Command safety",
                "Commands carry authority, expiry, idempotency and acknowledgement; stale commands are not executed after reconnect.",
            ),
            (
                "Graceful degradation",
                "Loss, delay, restart and jamming scenarios produce visible stale/degraded state, growing uncertainty and clean recovery.",
            ),
            (
                "Protocol restraint",
                "No additional protocol is accepted without named hardware, an owner, a fixture and an end-to-end acceptance test.",
            ),
        ],
        [1900, 7460],
    )
    add_callout(
        doc,
        "Bottom line",
        "Vehicle- and sensor-connected C2 is not a fairytale. The fairytale is assuming every asset is integrated, every link is continuous, and encrypted connectivity is tactically free. Sentinel should make uncertainty, provenance and network reality part of the product.",
        accent=GOLD,
        fill="FFF8E8",
    )

    source_section = doc.add_section(WD_SECTION.CONTINUOUS)
    configure_section(source_section, landscape=False, margin=1.0)
    source_section.different_first_page_header_footer = False
    source_section.header.is_linked_to_previous = True
    source_section.footer.is_linked_to_previous = True
    doc.add_heading("Appendix A. Source register", level=1)
    add_small_note(
        doc,
        "Public sources used for hardware handshakes, platform benchmarks, standards, and combat cases. Vendor claims are treated as descriptions of stated capability, not independent validation.",
        after=8,
    )
    sources = [
        (1, "PX4 - System architecture", "https://docs.px4.io/v1.14/en/concept/px4_systems_architecture", "Autopilot, telemetry and companion-computer responsibilities."),
        (2, "PX4 - Companion computers", "https://docs.px4.io/v1.14/en/companion_computer/", "Linux companion-computer role and vehicle connection."),
        (3, "MAVLink - Heartbeat protocol", "https://mavlink.io/en/services/heartbeat.html", "Vehicle/component discovery and connection state."),
        (4, "MAVLink - Component metadata", "https://mavlink.io/en/services/component_metadata.html", "Capability and metadata discovery."),
        (5, "MAVLink - Security FAQ", "https://mavlink.io/en/about/faq.html", "Message signing and the lack of payload encryption."),
        (6, "MAVLink - Command protocol", "https://mavlink.io/en/services/command.html", "Command acknowledgement and retry semantics."),
        (7, "Auterion - Skynode peripheral integration", "https://docs.auterion.com/hardware-integration/skynode/peripherals", "Mission computer, autopilot and peripheral interfaces."),
        (8, "Anduril - Lattice entities overview", "https://developer.anduril.com/guides/entities/overview", "Composable entity, provenance, expiry and partial-state model."),
        (9, "Anduril - Lattice SDK", "https://www.anduril.com/lattice/lattice-sdk", "REST/gRPC integration, tasks and edge object distribution."),
        (10, "Picogrid - Engineering and tactical-edge architecture", "https://picogrid.com/newsroom/engineering-at-picogrid", "Legion common interface and Lander/Helios edge compute."),
        (11, "UK Government - SAPIENT autonomous sensor system", "https://www.gov.uk/guidance/sapient-autonomous-sensor-system", "Open sensor/effector messaging and tasking."),
        (12, "ONVIF - Profile T", "https://www.onvif.org/profiles/profile-t/", "IP video streaming, metadata and control."),
        (13, "MITRE - Cursor on Target overview", "https://www.mitre.org/news-insights/publication/cursor-target", "Tactical event exchange background."),
        (14, "US Army - VICTORY open vehicle architecture", "https://www.army.mil/article/210117/reusable_and_refresh_able_open_systems_architecture_for_fighting_vehicles", "Vehicle data/service sharing above local buses."),
        (15, "US Army Armor - Blue Force Tracking in Operation Iraqi Freedom", "https://www.benning.army.mil/armor/eARMOR/content/issues/2003/SEP_OCT/ArmorSeptemberOctober2003web.pdf", "Field observations, scale, update behavior and constraints."),
        (16, "US GAO - Battlefield automation interoperability in OIF", "https://www.gao.gov/assets/gao-04-547.pdf", "Incompatible tracking systems and improvised interoperability."),
        (17, "US Army - Blue Force Tracking upgrade in Afghanistan", "https://www.army.mil/article/96438/army_upgrades_blue_force_tracking_in_afghanistan_to_prepare_for_new_network", "Vehicle and tactical-operations-centre kit rollout."),
        (18, "NATO ACT - DELTA at CWIX", "https://www.act.nato.int/article/delta-system-cwix/", "Interoperability and multi-source integration."),
        (19, "Ukraine Ministry of Defence - DELTA Mission Control adoption", "https://mod.gov.ua/en/news/kateryna-chernohorenko-the-mission-control-module-is-now-accessible-to-all-delta-users", "Published units, daily missions and cumulative mission figures."),
        (20, "RUSI - Preliminary lessons from Russia's war in Ukraine", "https://static.rusi.org/special-report-202207-ukraine-final-web_0.pdf", "Communications, commercial satellite dependence and operational lessons."),
        (21, "RUSI - Russian communications and the world of hertz", "https://www.rusi.org/explore-our-research/publications/commentary/russian-comms-ukraine-world-hertz", "Radio-frequency detection and geolocation risk."),
        (22, "Y Combinator - Surtr Defense Systems / ParallaxOS", "https://www.ycombinator.com/companies/surtr-defense-systems", "Vendor description of hardware-agnostic counter-UAS sensor fusion."),
        (23, "US Air Force Research Laboratory - TAK during Afghanistan evacuation", "https://www.afrl.af.mil/News/Article-Display/Article/3360125/afrl-technology-aids-operators-during-afghanistan-evacuation/", "Mission-specific digital coordination case."),
    ]
    for source in sources:
        add_source(doc, *source)

    props = doc.core_properties
    props.title = "Sentinel Network Architecture"
    props.subject = "MVP and demo network topology for sensors, vehicles, UAVs, simulation and edge integration"
    props.author = "Sentinel Project Team"
    props.keywords = "Sentinel, C2, edge gateway, MAVLink, sensors, Gazebo, network architecture"
    props.comments = "Architecture decision and evidence synthesis for Sentinel MVP planning."
    OUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT)
    print(OUT)
    print(DIAGRAM)


if __name__ == "__main__":
    build_doc()
